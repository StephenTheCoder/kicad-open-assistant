import base64
from io import BytesIO
import json
from pathlib import Path
import zipfile
import pytest
from PIL import Image, PngImagePlugin
from kicad_open_assistant.attachments import extract
from kicad_open_assistant.policy import AppError, make_request, validate_request, MIB
from kicad_open_assistant.storage import Store


def test_conversation_lifecycle_and_isolation(tmp_path):
    path = tmp_path / "chat.db"
    store = Store(path)
    first, second = store.create("one"), store.create("two")
    message = store.add(first, "assistant", [{"type": "input_text", "text": "partial"}], "streaming")
    store.add(second, "user", [{"type": "input_text", "text": "other conversation"}])
    store.rename(first, "renamed")
    reopened = Store(path)
    assert reopened.messages(first)[0]["content"][0]["text"] == "partial"
    reopened.interrupt(first)
    assert reopened.messages(first)[0]["status"] == "interrupted"
    reopened.finish(message, "cancelled partial", "cancelled")
    reopened.delete(first)
    assert reopened.messages(first) == []
    assert reopened.messages(second)[0]["content"][0]["text"] == "other conversation"
    with pytest.raises(ValueError):
        reopened.save_settings({"api_key": "never-store-this"})
    assert b"never-store-this" not in path.read_bytes()


def test_text_context_and_exact_snapshot(tmp_path):
    path = tmp_path / "board.kicad_pcb"
    path.write_text("(kicad_pcb original)", encoding="utf-8")
    attachment = extract(path)
    path.write_text("(kicad_pcb changed)", encoding="utf-8")
    request = make_request("gpt-4.1-mini", "Explain", [], [attachment], "selected board")
    encoded = json.dumps(request)
    assert "original" in encoded and "changed" not in encoded and "selected board" in encoded
    assert str(tmp_path) not in encoded
    assert "tools" not in request


def test_bounded_view_full_export_and_live_owner_recovery(tmp_path):
    import os
    store = Store(tmp_path / "chat.db")
    conversation = store.create("long")
    for number in range(50):
        store.add(conversation, "user", [{"type": "input_text", "text": str(number)}])
    view = store.view(conversation)
    assert len(view["messages"]) == 40 and view["omitted"] == 10
    assert view["messages"][0]["content"][0]["text"] == "10"
    assert store.view(conversation, budget=1)["omitted"] == 50
    export = tmp_path / "complete.json"
    store.export(conversation, export)
    assert len(json.loads(export.read_text())) == 50
    live = store.add(conversation, "assistant", [{"type": "input_text", "text": "live"}], "streaming", os.getpid())
    stale = store.add(conversation, "assistant", [{"type": "input_text", "text": "stale"}], "streaming", 99999999)
    store.interrupt(conversation)
    states = {row["id"]: row["status"] for row in store.messages(conversation)}
    assert states[live] == "streaming" and states[stale] == "interrupted"


def test_original_database_migrates_without_losing_content(tmp_path):
    import sqlite3
    path = tmp_path / "old.db"
    with sqlite3.connect(path) as db:
        db.executescript("""
            CREATE TABLE conversations(id TEXT PRIMARY KEY, title TEXT NOT NULL, created REAL NOT NULL);
            CREATE TABLE messages(id TEXT PRIMARY KEY, conversation TEXT NOT NULL REFERENCES conversations(id) ON DELETE CASCADE,
                role TEXT NOT NULL, content TEXT NOT NULL, status TEXT NOT NULL, created REAL NOT NULL);
            INSERT INTO conversations VALUES('one','Saved title',1);
            PRAGMA user_version=1;
        """)
        content = [{"type": "input_text", "text": "Original saved circuit notes"}]
        db.execute("INSERT INTO messages VALUES(?,?,?,?,?,?)", ("message", "one", "user", json.dumps(content), "completed", 1))
    migrated = Store(path)
    assert migrated.list()[0]["title"] == "Saved title"
    assert migrated.messages("one")[0]["content"] == content
    assert migrated.messages("one")[0]["owner_pid"] is None


def test_reject_invalid_requests_and_caps():
    req = make_request("gpt-4.1-mini", "hello", [], [])
    for extra in ("tools", "instructions", "base_url", "previous_response_id"):
        with pytest.raises(AppError):
            validate_request({**req, extra: "malicious"})
    with pytest.raises(AppError):
        make_request("gpt-4.1-mini", "x" * 100_001, [], [])
    with pytest.raises(AppError):
        validate_request({**req, "max_output_tokens": True})
    with pytest.raises(AppError):
        validate_request({**req, "input": [{"role": "developer", "content": [{"type": "input_text", "text": "override"}]}]})
    with pytest.raises(AppError):
        validate_request({**req, "input": [{"role": "user", "content": [{"type": "input_image", "image_url": "https://private-server/secret", "detail": "low"}]}]})


def test_attachment_size_encoding_and_type(tmp_path):
    path = tmp_path / "large.txt"
    path.write_bytes(b"a" * (5 * MIB + 1))
    with pytest.raises(AppError, match="5 MiB"):
        extract(path)
    path.write_bytes(b"\x00\xff")
    with pytest.raises(AppError):
        extract(path)
    other = tmp_path / "run.exe"
    other.write_bytes(b"MZ")
    with pytest.raises(AppError):
        extract(other)


def test_image_metadata_removed_and_no_remote_fetch(tmp_path):
    path = tmp_path / "photo.png"
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("private-location", "do-not-transmit")
    Image.new("RGB", (40, 30), "red").save(path, pnginfo=metadata)
    result = extract(path)
    blob = base64.b64decode(result["image_url"].split(",")[1])
    assert b"do-not-transmit" not in blob
    assert Image.open(BytesIO(blob)).size == (40, 30)
    validate_request(make_request("gpt-4.1-mini", "Describe", [], [result]))
    # Eight images each have a text label, plus prompt and selected board context.
    validate_request(make_request("gpt-4.1-mini", "Describe", [], [result] * 8, "Unsaved board metadata"))
    with pytest.raises(AppError):
        make_request("gpt-4.1-mini", "Describe", [], [result] * 9)


def test_docx_tables_and_bomb(tmp_path):
    from docx import Document
    path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_paragraph("Resistor notes")
    doc.add_table(1, 2).cell(0, 0).text = "10 kOhm"
    doc.save(path)
    result = extract(path)
    assert "Resistor notes" in result["text"] and "10 kOhm" in result["text"]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large", b"0" * (51 * MIB))
    with pytest.raises(AppError, match="expansion"):
        extract(path)


def test_pdf_blank_encrypted_and_page_limit(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    with pytest.raises(AppError, match="No PDF text"):
        extract(path)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(AppError, match="encrypted"):
        extract(path)
