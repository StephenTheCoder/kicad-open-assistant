"""Bounded local extraction. Called only in the disposable, limited worker."""
import base64
import hashlib
from io import BytesIO
from pathlib import Path
import warnings
import zipfile
from .policy import AppError, MAX_TEXT, MIB

TEXT = {".kicad_pcb", ".kicad_sch", ".kicad_pro", ".net", ".kicad_sym", ".kicad_mod", ".txt", ".md", ".csv", ".json"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp"}


def extract(path):
    source = Path(path)
    if source.is_symlink() or not source.is_file():
        raise AppError("file", "Select a regular file, not a directory or symbolic link.")
    suffix = source.suffix.lower()
    if suffix not in TEXT | IMAGES | {".pdf", ".docx"}:
        raise AppError("file_type", "Unsupported file. Use KiCad/text, PDF, DOCX, PNG, JPEG or WebP.")
    cap = 5 * MIB if suffix in TEXT else 10 * MIB
    if source.stat().st_size > cap:
        raise AppError("file_size", f"File exceeds {cap // MIB} MiB.")
    with source.open("rb") as stream:
        blob = stream.read(cap + 1)
    if len(blob) > cap:
        raise AppError("file_size", "File grew beyond its size limit.")
    result = {"name": source.name, "original_bytes": len(blob), "source_sha256": hashlib.sha256(blob).hexdigest()}
    try:
        if suffix in TEXT:
            text = blob.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError()
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(blob), strict=True)
            if reader.is_encrypted or len(reader.pages) > 100:
                raise AppError("pdf", "PDF is encrypted or exceeds 100 pages.")
            pages, count = [], 0
            for page in reader.pages:
                content = page.extract_text() or ""
                count += len(content)
                if count > MAX_TEXT:
                    raise AppError("text_limit", "PDF text exceeds 100,000 characters.")
                pages.append(content)
            text = "\n\n".join(pages)
            if any(not p.strip() for p in pages):
                result["warning"] = "Some PDF pages contain no extractable text; scanned images are not sent."
            if not text.strip():
                raise AppError("pdf_scanned", "No PDF text found. Attach selected page images instead; OCR is not included.")
        elif suffix == ".docx":
            with zipfile.ZipFile(BytesIO(blob)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(m.file_size for m in members) > 50 * MIB or any(m.flag_bits & 1 for m in members):
                    raise AppError("docx", "DOCX archive is encrypted or exceeds expansion limits.")
                if any("vbaproject" in m.filename.lower() for m in members):
                    raise AppError("docx", "Macro content is not accepted.")
            from docx import Document
            doc = Document(BytesIO(blob))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
            result["warning"] = "DOCX text and tables only; embedded pictures and charts are not transmitted."
        else:
            from PIL import Image, ImageOps
            with warnings.catch_warnings():
                warnings.simplefilter("error")
                with Image.open(BytesIO(blob)) as original:
                    if original.format not in {"PNG", "JPEG", "WEBP"} or original.width * original.height > 20_000_000 or getattr(original, "n_frames", 1) != 1:
                        raise AppError("image", "Use a single image of at most 20 megapixels.")
                    original.load()
                    im = ImageOps.exif_transpose(original).convert("RGB")
                    clean = Image.new("RGB", im.size)
                    clean.paste(im)
                    output = BytesIO()
                    clean.save(output, format="PNG")
                    outgoing = output.getvalue()
            if len(outgoing) > 10 * MIB:
                raise AppError("image", "Metadata-free PNG exceeds 10 MiB; resize the image before attaching.")
            result.update(image_url="data:image/png;base64," + base64.b64encode(outgoing).decode("ascii"), dimensions=list(clean.size))
            text = "Metadata-free PNG, low-detail vision input."
    except AppError:
        raise
    except Exception:
        raise AppError("file_invalid", "File is corrupt, unreadable, has unsupported encoding, or cannot be safely extracted.") from None
    if len(text) > MAX_TEXT:
        raise AppError("text_limit", "Extracted text exceeds 100,000 characters; select a smaller excerpt.")
    result["text"] = text
    return result
